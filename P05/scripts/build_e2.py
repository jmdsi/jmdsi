# -*- coding: utf-8 -*-
"""E2 — Expediente de Ingeniería Básica P05 en el formato exacto del P04-R02.
Genera HTML (plantilla/tipografía P04), PDF (A4 reportlab) y DOCX (Segoe UI)."""
import base64, glob, html as H, os, re
from common import COD, FECHA, CONF, SELLO, F_FICHA, F_CAT, F_P04, F_PM
from contenido import CAP

BASE = os.path.dirname(os.path.abspath(__file__))
ENT = os.path.join(BASE, "..", "entregables")
FIGD = os.path.join(ENT, "figuras")
LAMD = os.path.join(ENT, "laminas")
STEM = COD + "_EXPEDIENTE-INGENIERIA-BASICA_PATIO-CUBIERTA-DE-HONOR"

AZUL, ACERO, ROJO, GRIS, LINEA = "#0b3c5d", "#1d5b7a", "#a4262c", "#f4f6f8", "#c9d4dc"

FIGT = {
 "G-01": "Implantación general de la Proa Central (esquema 1:400)",
 "G-02": "Línea de tiempo patrimonial 1951–2030",
 "G-03": "Conciliación de áreas anidadas 1.625 / 380 / 210 / 1.245 m²",
 "G-04": "Retícula de paños y módulos de 4,00 m",
 "G-05": "Zonificación funcional ceremonial",
 "G-06": "Sección constructiva del pavimento: existente vs. propuesto",
 "G-07": "Mapa de daño post-sísmico D2 / D3 / D4",
 "G-08": "Esquema de drenaje pluvial (Tr = 25 años)",
 "G-09": "Unifilar eléctrico y protección atmosférica",
 "G-10": "Emplazamiento del Mástil OMM y elementos verticales",
 "G-11": "Estrategia HBIM y gemelo digital",
 "G-12": "Plano de evacuación R-1 y aforos",
 "G-13": "ROM en cascada (895.000 → 565.000 USD)",
 "G-14": "Cronograma de fases con gates",
 "G-15": "Mapa de riesgos (probabilidad × impacto)",
}

def figpng(g):
    n = g.replace("G-", "G")
    m = glob.glob(os.path.join(FIGD, "*FIGURA-%s_*.png" % n))
    return m[0] if m else None

COVER_ROWS = [
 ("Código del expediente", COD),
 ("Edificaciones", "Patio / Cubierta de Honor (EXT-PCH-04 · UMC-EXT-004 / UMC-ED-000-PCH04) — explanada ceremonial de la Proa Central"),
 ("Proyecto", "Proyecto 05 del Plan Maestro UMC 2025–2030 · Espacios Exteriores Institucionales / Área Cultural y Ceremonial"),
 ("Fase", "Ingeniería Básica / Anteproyecto"),
 ("Revisión y fecha", "R01 · " + FECHA),
 ("Emplazamiento", "Campus UMC, Catia La Mar, parroquia Catia La Mar, estado La Guaira, República Bolivariana de Venezuela"),
 ("Sistema de referencia", "WGS84 / EPSG:4326 (trabajo UTM 19N · EPSG:32619)"),
 ("Clase de estimación", "AACE Clase 5 (ROM) · exactitud −25 % / +40 %"),
 ("Evento condicionante", "Doblete sísmico del 24-06-2026 (Mw 7,2 y Mw 7,5) · patio INHABILITADO como punto R-1"),
 ("Aprobación", "Rectorado UMC – Cap. Alt. / Dr. Guillermo Antonio Riut Hernández"),
 ("Autoría del documento fuente", "Prof. Jhonny Da Silva Izquierdo (MSc. Ing. Lic., CIV 95787)"),
 ("Proyectista responsable", "Por designar – profesional habilitado"),
 ("Confidencialidad", CONF),
 ("Estado de emisión", "Emitido para revisión técnica, coordinación interdisciplinaria y presupuesto preliminar"),
]

CSS = """
:root{--azul:#0b3c5d;--acero:#1d5b7a;--rojo:#a4262c;--gris:#f4f6f8;--linea:#c9d4dc}
*{box-sizing:border-box}body{font-family:"Segoe UI",Helvetica,Arial,sans-serif;color:#12212b;
 line-height:1.55;margin:0;background:#e8ecef}
nav#toc{position:fixed;left:0;top:0;width:250px;height:100vh;overflow:auto;background:var(--azul);
 color:#dce6ee;padding:18px 14px;font-size:12px}
nav#toc a{color:#cfe0ea;text-decoration:none;display:block;padding:2px 0}
nav#toc a.h2{padding-left:12px;font-size:11px;color:#9db8c9}
main{margin-left:250px;max-width:900px;background:#fff;padding:40px 48px 80px;
 box-shadow:0 2px 16px rgba(0,0,0,.15)}
h1{color:var(--azul);font-size:1.45rem;border-bottom:3px solid var(--azul);padding-bottom:.3em;
 margin-top:2.2em;page-break-after:avoid}
h2{color:var(--azul);font-size:1.12rem;margin-top:1.7em;page-break-after:avoid}
h3{color:var(--acero);font-size:1rem;margin-top:1.4em}
p,li{font-size:.9rem}
table{border-collapse:collapse;width:100%;margin:1em 0;font-size:.8rem;page-break-inside:auto}
th{background:var(--azul);color:#fff;text-align:left;padding:6px 8px;border:1px solid var(--azul)}
td{border:1px solid var(--linea);padding:5px 8px;vertical-align:top}
tbody tr:nth-child(even){background:var(--gris)}
tr{page-break-inside:avoid}
figure{margin:1.2em 0;text-align:center;page-break-inside:avoid}
figure img{max-width:100%;border:1px solid var(--linea)}
figcaption{font-size:.76rem;color:#33475a;margin-top:.4em}
.note{border-left:5px solid var(--rojo);background:#fdf3f3;padding:.7em 1em;margin:1.1em 0;
 font-size:.85rem;white-space:pre-line}
.cover{background:var(--azul);color:#fff;padding:44px 40px;margin:-40px -48px 30px}
.cover h1{color:#fff;border:0;margin:0 0 .3em;font-size:1.9rem}
.cover table{color:#12212b}
.sello{color:#ffd9d6;font-weight:700;font-size:.8rem;margin-top:1.4em}
@media print{@page{size:A4;margin:15mm 13mm}nav#toc{display:none}main{margin:0;box-shadow:none;
 padding:0;max-width:none}.cover{margin:0 0 20mm}}
"""

def esc(t):
    return H.escape(t, quote=False)

def build_html():
    parts, nav = [], []
    nav.append("<a class='h1' href='#0'>ÍNDICE DEL EXPEDIENTE</a>")
    # cover
    cov = ["<main><div class='cover'><h1>Proyecto de Ingeniería Básica — Patio / Cubierta de Honor "
           "(«La Cubierta de Formación»)</h1><table>"]
    for k, v in COVER_ROWS:
        cov.append("<tr><th style='width:34%%'>%s</th><td>%s</td></tr>" % (esc(k), esc(v)))
    cov.append("</table>\n<div class='sello'>%s</div></div>" % esc(SELLO))
    parts.append("".join(cov))
    # índice
    parts.append("<h1 id='0'>ÍNDICE DEL EXPEDIENTE</h1><p>Veintiocho capítulos conforme al esquema "
                 "rector del Plan Maestro UMC 2025–2030 (expediente P01) y al formato del expediente "
                 "P04-R02, adaptados al objeto de este proyecto.</p>")
    idx = ["<table><thead><tr><th>Capítulos 1 – 14</th><th>Capítulos 15 – 28</th></tr></thead><tbody>"]
    for i in range(14):
        a = "%d. %s" % (CAP[i][0], CAP[i][1])
        b = "%d. %s" % (CAP[i + 14][0], CAP[i + 14][1])
        idx.append("<tr><td>%s</td><td>%s</td></tr>" % (esc(a), esc(b)))
    idx.append("</tbody></table>")
    parts.append("".join(idx))
    # capítulos
    for n, tit, blo in CAP:
        nav.append("<a class='h1' href='#%d'>%d. %s</a>" % (n, n, esc(tit)))
        parts.append("<h1 id='%d'>%d. %s</h1>" % (n, n, esc(tit)))
        j = 0
        for b in blo:
            if b[0] == "p":
                parts.append("<p>%s</p>" % esc(b[1]))
            elif b[0] == "h3":
                j += 1
                hid = "%d-%d" % (n, j)
                nav.append("<a class='h2' href='#%s'>%s</a>" % (hid, esc(b[1])))
                parts.append("<h3 id='%s'>%s</h3>" % (hid, esc(b[1])))
            elif b[0] == "box":
                parts.append("<div class='note'>%s</div>" % esc(b[1]))
            elif b[0] == "fig":
                g = b[1]; p = figpng(g)
                if p:
                    d = base64.b64encode(open(p, "rb").read()).decode()
                    parts.append("<figure><img src='data:image/png;base64,%s' alt='%s'>"
                                 "<figcaption>Figura %s — %s · %s · %s</figcaption></figure>"
                                 % (d, g, g, esc(FIGT[g]), COD, "NO APTO PARA CONSTRUCCIÓN"))
            elif b[0] == "t":
                enc, filas, pie = b[1], b[2], b[3]
                t = ["<table><thead><tr>"] + ["<th>%s</th>" % esc(c) for c in enc] + ["</tr></thead><tbody>"]
                for f in filas:
                    t.append("<tr>" + "".join("<td>%s</td>" % esc(c) for c in f) + "</tr>")
                t.append("</tbody></table>")
                parts.append("".join(t))
                if pie:
                    parts.append("<p style='font-size:.78rem;color:#33475a'><em>%s</em></p>" % esc(pie))
    # pie documental
    parts.append("<hr><p style='font-size:.76rem;color:#33475a'>%s · %s · %s · Documento rector: %s. "
                 "Sello de emisión: %s.</p></main>" % (COD, esc(FECHA), esc(CONF), esc(F_FICHA), esc(SELLO)))
    doc = ("<!DOCTYPE html><html lang='es'><head><meta charset='utf-8'>"
           "<meta name='viewport' content='width=device-width,initial-scale=1'>"
           "<title>%s — Expediente de Ingeniería Básica</title><style>%s</style></head><body>"
           "<nav id='toc'><strong>%s</strong><hr style='border-color:#2c5a7c'>%s</nav>%s</body></html>"
           % (COD, CSS, COD, "".join(nav), "".join(parts)))
    out = os.path.join(ENT, STEM + ".html")
    open(out, "w", encoding="utf-8").write(doc)
    print("HTML", out, os.path.getsize(out))

# ---------------------------------------------------------------- PDF
def build_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
                                    Table, TableStyle, Image, PageBreak, NextPageTemplate)
    cAzul, cAcero, cRojo = colors.HexColor(AZUL), colors.HexColor(ACERO), colors.HexColor(ROJO)
    cGris, cLinea = colors.HexColor(GRIS), colors.HexColor(LINEA)
    W, Hh = A4
    ML, MR, MT, MB = 13 * mm, 13 * mm, 15 * mm, 17 * mm
    CW = W - ML - MR

    S = lambda **k: ParagraphStyle(k.pop("name"), **k)
    st_p = S(name="p", fontName="Helvetica", fontSize=9, leading=13.2, spaceAfter=5)
    st_h1 = S(name="h1", fontName="Helvetica-Bold", fontSize=14.5, leading=18, textColor=cAzul,
              spaceBefore=14, spaceAfter=8)
    st_h3 = S(name="h3", fontName="Helvetica-Bold", fontSize=10.5, leading=14, textColor=cAcero,
              spaceBefore=9, spaceAfter=4)
    st_cell = S(name="c", fontName="Helvetica", fontSize=7.3, leading=9.4)
    st_cellh = S(name="ch", fontName="Helvetica-Bold", fontSize=7.5, leading=9.6, textColor=colors.white)
    st_note = S(name="n", fontName="Helvetica", fontSize=8.6, leading=12, textColor=colors.HexColor("#5a1518"))
    st_pie = S(name="pie", fontName="Helvetica-Oblique", fontSize=7.6, leading=10,
               textColor=colors.HexColor("#33475a"), spaceAfter=6)
    st_cap = S(name="cap", fontName="Helvetica", fontSize=7.6, leading=10,
               textColor=colors.HexColor("#33475a"), alignment=1, spaceAfter=8)

    def hr_h1(txt):
        t = Table([[Paragraph(txt, st_h1)]], colWidths=[CW])
        t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 2.2, cAzul),
                               ("LEFTPADDING", (0, 0), (-1, -1), 0),
                               ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        return t

    def note(txt):
        t = Table([[Paragraph(txt.replace("\n", "<br/>"), st_note)]], colWidths=[CW])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fdf3f3")),
                               ("LINEBEFORE", (0, 0), (0, -1), 3.4, cRojo),
                               ("LEFTPADDING", (0, 0), (-1, -1), 9),
                               ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                               ("TOPPADDING", (0, 0), (-1, -1), 6),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        return t

    def tabla(enc, filas, pie):
        data = [[Paragraph(c, st_cellh) for c in enc]]
        for f in filas:
            data.append([Paragraph(c, st_cell) for c in f])
        nc = len(enc)
        if nc == 2:
            cw = [CW * .3, CW * .7]
        elif nc == 3:
            cw = [CW * .3, CW * .4, CW * .3]
        else:
            cw = [CW / nc] * nc
        t = Table(data, colWidths=cw, repeatRows=1)
        sty = [("BACKGROUND", (0, 0), (-1, 0), cAzul),
               ("GRID", (0, 0), (-1, -1), .5, cLinea),
               ("BOX", (0, 0), (-1, 0), .5, cAzul),
               ("VALIGN", (0, 0), (-1, -1), "TOP"),
               ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
               ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        for i in range(1, len(data)):
            if i % 2 == 0:
                sty.append(("BACKGROUND", (0, i), (-1, i), cGris))
        t.setStyle(TableStyle(sty))
        out = [t]
        if pie:
            out.append(Paragraph(pie, st_pie))
        return out

    def fig(g):
        p = figpng(g)
        if not p:
            return []
        from PIL import Image as PImage
        iw, ih = PImage.open(p).size
        w = min(CW, 158 * mm)
        h = w * ih / iw
        maxh = 165 * mm
        if h > maxh:
            w, h = maxh * iw / ih, maxh
        return [Spacer(1, 4),
                Image(p, width=w, height=h),
                Paragraph("Figura %s — %s · %s · NO APTO PARA CONSTRUCCIÓN" % (g, FIGT[g], COD), st_cap)]

    story = []
    # portada
    st_cvt = S(name="cvt", fontName="Helvetica-Bold", fontSize=17, leading=22, textColor=colors.white)
    st_cvk = S(name="cvk", fontName="Helvetica-Bold", fontSize=7.8, leading=10, textColor=colors.white)
    st_cvv = S(name="cvv", fontName="Helvetica", fontSize=7.8, leading=10, textColor=colors.HexColor("#dce6ee"))
    st_cvs = S(name="cvs", fontName="Helvetica-Bold", fontSize=8.2, leading=11, textColor=colors.HexColor("#ffd9d6"))
    rows = [[Paragraph("Proyecto de Ingeniería Básica — Patio / Cubierta de Honor («La Cubierta de Formación»)", st_cvt), ""]]
    for k, v in COVER_ROWS:
        rows.append([Paragraph(k, st_cvk), Paragraph(v, st_cvv)])
    rows.append([Paragraph(SELLO, st_cvs), ""])
    cov = Table(rows, colWidths=[CW * .34, CW * .66])
    cs = [("BACKGROUND", (0, 0), (-1, -1), cAzul),
          ("SPAN", (0, 0), (1, 0)), ("SPAN", (0, len(rows) - 1), (1, len(rows) - 1)),
          ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
          ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
          ("TOPPADDING", (0, 0), (1, 0), 16), ("BOTTOMPADDING", (0, 0), (1, 0), 12),
          ("TOPPADDING", (0, len(rows) - 1), (1, len(rows) - 1), 12),
          ("BOTTOMPADDING", (0, len(rows) - 1), (1, len(rows) - 1), 16)]
    for i in range(1, len(rows) - 1):
        cs.append(("LINEBELOW", (0, i), (1, i), .4, colors.HexColor("#1d5b7a")))
    cov.setStyle(TableStyle(cs))
    story += [cov, PageBreak()]
    # índice
    story.append(hr_h1("ÍNDICE DEL EXPEDIENTE"))
    story.append(Paragraph("Veintiocho capítulos conforme al esquema rector del Plan Maestro UMC "
                           "2025–2030 (expediente P01) y al formato del expediente P04-R02.", st_p))
    story += tabla(["Capítulos 1 – 14", "Capítulos 15 – 28"],
                   [["%d. %s" % (CAP[i][0], CAP[i][1]), "%d. %s" % (CAP[i + 14][0], CAP[i + 14][1])]
                    for i in range(14)], "")
    for n, tit, blo in CAP:
        story.append(PageBreak())
        story.append(hr_h1("%d. %s" % (n, tit)))
        for b in blo:
            if b[0] == "p":
                story.append(Paragraph(b[1], st_p))
            elif b[0] == "h3":
                story.append(Paragraph(b[1], st_h3))
            elif b[0] == "box":
                story.append(note(b[1])); story.append(Spacer(1, 5))
            elif b[0] == "t":
                story += tabla(b[1], b[2], b[3])
            elif b[0] == "fig":
                story += fig(b[1])
    # Anexo P: láminas apaisadas
    story.append(NextPageTemplate("lamina"))
    story.append(PageBreak())
    st_lt = S(name="lt", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=cAzul, spaceAfter=4)
    lams = sorted(glob.glob(os.path.join(LAMD, "*LAMINA-P*.png")))
    LW, LH = landscape(A4)
    lcw, lch = LW - 20 * mm, LH - 34 * mm
    from PIL import Image as PImage
    for i, p in enumerate(lams):
        name = os.path.basename(p)
        m = re.search(r"LAMINA-(P\d)_(.+)\.png", name)
        lam = m.group(1)
        tit = m.group(2).replace("-", " ")
        iw, ih = PImage.open(p).size
        w = lcw; h = w * ih / iw
        if h > lch:
            w, h = lch * iw / ih, lch
        story.append(Paragraph("ANEXO P · LÁMINA %s — %s" % (lam, tit), st_lt))
        story.append(Image(p, width=w, height=h))
        if i < len(lams) - 1:
            story.append(PageBreak())

    out = os.path.join(ENT, STEM + ".pdf")
    doc = BaseDocTemplate(out, pagesize=A4, leftMargin=ML, rightMargin=MR,
                          topMargin=MT, bottomMargin=MB, title=COD + " Expediente de Ingeniería Básica",
                          author="UMC — Plan Maestro 2025-2030")

    def decor(canv, d, land=False):
        pw = landscape(A4)[0] if land else W
        canv.saveState()
        canv.setStrokeColor(cAzul); canv.setLineWidth(.6)
        canv.setFont("Helvetica", 6.4); canv.setFillColor(colors.HexColor("#33475a"))
        canv.drawString(ML, 9.5 * mm, COD + " · Expediente de Ingeniería Básica · " + FECHA)
        canv.setFillColor(cRojo); canv.setFont("Helvetica-Bold", 6.4)
        canv.drawCentredString(pw / 2, 5.5 * mm, "NO APTO PARA CONSTRUCCIÓN — AACE CLASE 5 (−25 % / +40 %)")
        canv.setFillColor(colors.HexColor("#33475a")); canv.setFont("Helvetica", 6.4)
        canv.drawRightString(pw - MR, 9.5 * mm, "Pág. %d" % canv.getPageNumber())
        canv.restoreState()

    fr_p = Frame(ML, MB, W - ML - MR, Hh - MT - MB, id="fp")
    fr_l = Frame(10 * mm, 14 * mm, LW - 20 * mm, LH - 26 * mm, id="fl")
    doc.addPageTemplates([
        PageTemplate(id="portrait", frames=[fr_p], onPage=lambda c, d: decor(c, d, False)),
        PageTemplate(id="lamina", frames=[fr_l], onPage=lambda c, d: decor(c, d, True),
                     pagesize=landscape(A4)),
    ])
    doc.build(story)
    data = open(out, "rb").read()
    print("PDF", out, os.path.getsize(out), "págs:",
          data.count(b"/Type /Page") - data.count(b"/Type /Pages"))

# ---------------------------------------------------------------- DOCX
def build_docx():
    import docx
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    AZ = RGBColor(0x0b, 0x3c, 0x5d); AC = RGBColor(0x1d, 0x5b, 0x7a); RJ = RGBColor(0xa4, 0x26, 0x2c)
    d = docx.Document()
    for sec in d.sections:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(13)
        sec.top_margin, sec.bottom_margin = Mm(15), Mm(15)
    base = d.styles["Normal"]
    base.font.name = "Segoe UI"; base.font.size = Pt(9)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")

    def shade(cell, hexcol):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexcol)
        tcPr.append(sh)

    def h1(t):
        p = d.add_paragraph(); r = p.add_run(t)
        r.bold = True; r.font.size = Pt(14); r.font.color.rgb = AZ; r.font.name = "Segoe UI"
        p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)

    def h3(t):
        p = d.add_paragraph(); r = p.add_run(t)
        r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = AC; r.font.name = "Segoe UI"

    def para(t, size=9, italic=False, color=None):
        p = d.add_paragraph(); r = p.add_run(t)
        r.font.size = Pt(size); r.font.name = "Segoe UI"; r.italic = italic
        if color: r.font.color.rgb = color
        return p

    def box(t):
        tb = d.add_table(rows=1, cols=1); tb.style = "Table Grid"
        c = tb.rows[0].cells[0]; shade(c, "fdf3f3")
        r = c.paragraphs[0].add_run(t)
        r.font.size = Pt(8.5); r.font.name = "Segoe UI"; r.font.color.rgb = RJ; r.bold = True

    def tabla(enc, filas, pie):
        tb = d.add_table(rows=1 + len(filas), cols=len(enc)); tb.style = "Table Grid"
        for j, c in enumerate(enc):
            cell = tb.rows[0].cells[j]; shade(cell, "0b3c5d")
            r = cell.paragraphs[0].add_run(c)
            r.bold = True; r.font.size = Pt(7.5); r.font.name = "Segoe UI"
            r.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        for i, f in enumerate(filas):
            for j, c in enumerate(f):
                cell = tb.rows[i + 1].cells[j]
                if i % 2 == 1: shade(cell, "f4f6f8")
                r = cell.paragraphs[0].add_run(c)
                r.font.size = Pt(7.3); r.font.name = "Segoe UI"
        if pie:
            para(pie, 7.6, True)

    # portada
    cv = d.add_table(rows=2 + len(COVER_ROWS), cols=2); cv.style = "Table Grid"
    tcell = cv.rows[0].cells[0].merge(cv.rows[0].cells[1]); shade(tcell, "0b3c5d")
    r = tcell.paragraphs[0].add_run("Proyecto de Ingeniería Básica — Patio / Cubierta de Honor "
                                    "(«La Cubierta de Formación»)")
    r.bold = True; r.font.size = Pt(16); r.font.name = "Segoe UI"; r.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
    for i, (k, v) in enumerate(COVER_ROWS):
        ck, cvv = cv.rows[i + 1].cells
        shade(ck, "0b3c5d"); shade(cvv, "12476b")
        rk = ck.paragraphs[0].add_run(k); rk.bold = True; rk.font.size = Pt(7.8)
        rk.font.name = "Segoe UI"; rk.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        rv = cvv.paragraphs[0].add_run(v); rv.font.size = Pt(7.8)
        rv.font.name = "Segoe UI"; rv.font.color.rgb = RGBColor(0xdc, 0xe6, 0xee)
    scell = cv.rows[-1].cells[0].merge(cv.rows[-1].cells[1]); shade(scell, "0b3c5d")
    rs = scell.paragraphs[0].add_run(SELLO)
    rs.bold = True; rs.font.size = Pt(8.5); rs.font.name = "Segoe UI"
    rs.font.color.rgb = RGBColor(0xff, 0xd9, 0xd6)
    d.add_page_break()
    # índice
    h1("ÍNDICE DEL EXPEDIENTE")
    tabla(["Capítulos 1 – 14", "Capítulos 15 – 28"],
          [["%d. %s" % (CAP[i][0], CAP[i][1]), "%d. %s" % (CAP[i + 14][0], CAP[i + 14][1])]
           for i in range(14)], "")
    for n, tit, blo in CAP:
        d.add_page_break()
        h1("%d. %s" % (n, tit))
        for b in blo:
            if b[0] == "p": para(b[1])
            elif b[0] == "h3": h3(b[1])
            elif b[0] == "box": box(b[1]); para("", 4)
            elif b[0] == "t": tabla(b[1], b[2], b[3])
            elif b[0] == "fig":
                p = figpng(b[1])
                if p:
                    pp = d.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pp.add_run().add_picture(p, width=Mm(170))
                    cap = para("Figura %s — %s · %s · NO APTO PARA CONSTRUCCIÓN"
                               % (b[1], FIGT[b[1]], COD), 7.6, True)
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para("%s · %s · %s · Documento rector: %s · %s" % (COD, FECHA, CONF, F_FICHA, SELLO), 7.6, True)
    out = os.path.join(ENT, STEM + ".docx")
    d.save(out)
    print("DOCX", out, os.path.getsize(out))

if __name__ == "__main__":
    build_html()
    build_pdf()
    build_docx()
