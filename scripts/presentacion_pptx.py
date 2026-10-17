#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""E6 v2 · Presentación para el Consejo Universitario con el registro fotográfico institucional
aportado por el Rectorado (4 imágenes): generador PPTX (python-pptx) + PDF + HTML.
28 láminas (rango 24–28 del prompt). Reutiliza el guion de presentacion.py."""
import copy
import os

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

import presentacion as PR

AZUL = RGBColor(0x0B, 0x3C, 0x5D)
ACERO = RGBColor(0x1D, 0x5B, 0x7A)
ROJO = RGBColor(0xA4, 0x26, 0x2C)
GRIS = RGBColor(0x33, 0x47, 0x5A)
CLARO = RGBColor(0xEE, 0xF2, 0xF5)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONDO_NOTA = RGBColor(0xFD, 0xF3, 0xF3)
LINEA = RGBColor(0xC9, 0xD4, 0xDC)

A = "src/assets/"
F1 = A + "IMG-CU-01_TORRE-ILUMINADA-NOCHE.png"
F2 = A + "IMG-CU-02_PLAZA-DE-HONOR-CREPUSCULO.png"
F3 = A + "IMG-CU-03_EVACUACION-DOBLETE-24JUN2026.png"
F4 = A + "IMG-CU-04_ESCUELA-NAUTICA-MALAUSSENA-HISTORICA.png"
G02 = "build/fig/G-02_linea_tiempo.png"

COD = PR.COD
SELLO = PR.SELLO


def deck():
    """Guion modificado: fotografías institucionales integradas y lámina del emblema."""
    S = copy.deepcopy(PR.S)
    out, inserted = [], False
    for sl in S:
        t = sl["t"]
        if t.startswith("Marco:"):
            sl["img"] = G02
        elif t.startswith("Patrimonio:"):
            sl["img"] = F4
            sl["nota"] = (sl.get("nota", "") + "  Fotografía histórica: archivo MOP, Escuela "
                          "Náutica de Venezuela, Luis Malaussena, c. 1951-1958 (registro aportado "
                          "por el Rectorado).").strip()
        elif t.startswith("El doblete sísmico"):
            sl["img"] = F3
            sl["nota"] = ("Fotograma del registro de evacuación del 24-06-2026 aportado por la "
                          "comunidad universitaria: se incorpora como evidencia referencial del "
                          "evento (anexo D y entregable E5). La línea de tiempo documental G-02 "
                          "queda en la lámina 4 y en el expediente.")
        elif t.startswith("Implantación del Área Cultural"):
            out.append(sl)
            out.append(dict(
                t="El emblema vigente: la estrella del vigía",
                st="La Torre iluminada sobre la plaza de honor: el valor ceremonial que el "
                   "proyecto protege y restituye",
                bullets=[
                    "La imagen nocturna del conjunto —Auditorio curvo, galería de protocolo y "
                    "Torre con el emblema encendido— es la síntesis del valor que se restaura.",
                    "El balizamiento escénico existente se conserva y se recablea: forma parte "
                    "del uso ceremonial documentado (actos, grados y fiestas patrias).",
                    "La plaza de honor opera como explanada de protocolo y como punto de "
                    "reunión R-1 del plan de evacuación del Área Cultural.",
                    "Contraste documental: la misma composición aparece en la fotografía "
                    "histórica de Malaussena (c. 1951-1958), con la torre original de 12,20 m."],
                img=F1, img2=F2,
                nota="Tomas nocturnas del conjunto Auditorio–Torre con el emblema institucional "
                     "iluminado y crepúsculo sobre la plaza de honor. Registro fotográfico "
                     "institucional UMC, 2026 (aportado por el Rectorado)."))
            inserted = True
            continue
        out.append(sl)
    assert inserted
    return out


# ------------------------------------------------------------------ PPTX
def _txt(slide, l, t, w, h, runs, size=14, color=GRIS, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=6, line=1.15):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [runs]
    for i, r in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line
        run = p.add_run()
        run.text = r
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = "Calibri"
    return box


def _rect(slide, l, t, w, h, fill, line_color=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(.75)
    sp.shadow.inherit = False
    return sp


def _fit(path, bw, bh):
    iw, ih = Image.open(path).size
    sc = min(bw / iw, bh / ih)
    return iw * sc, ih * sc


def _img(slide, path, l, t, bw, bh):
    if not os.path.exists(path):
        return
    w, h = _fit(path, bw, bh)
    slide.shapes.add_picture(path, Inches(l + (bw - w) / 2), Inches(t + (bh - h) / 2),
                             Inches(w), Inches(h))


def render_pptx(S, path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    n = len(S)
    for i, sl in enumerate(S, 1):
        s = prs.slides.add_slide(blank)
        # banda de título
        _rect(s, 0, 0, 13.333, 1.32, AZUL)
        _rect(s, 0, 1.32, 13.333, 0.075, ROJO)
        _txt(s, 0.45, 0.22, 12.45, 0.75, sl["t"], size=25, color=BLANCO, bold=True, space_after=2)
        _txt(s, 0.47, 0.86, 12.4, 0.4, sl["st"], size=12, color=RGBColor(0xCF, 0xE0, 0xEA),
             space_after=0)
        y = 1.62
        left_w = 12.45
        imgs = [x for x in (sl.get("img"), sl.get("img2")) if x and os.path.exists(x)]
        if imgs:
            left_w = 6.9
        # KPIs
        if sl.get("kpis"):
            ks = sl["kpis"]
            kw = left_w / len(ks)
            for j, (v, l) in enumerate(ks):
                x = 0.45 + j * kw
                _rect(s, x + 0.06, y, kw - 0.16, 0.92, CLARO)
                _rect(s, x + 0.06, y, kw - 0.16, 0.06, ACERO)
                _txt(s, x + 0.12, y + 0.12, kw - 0.28, 0.45, v, size=19, color=AZUL, bold=True,
                     align=PP_ALIGN.CENTER, space_after=0)
                _txt(s, x + 0.12, y + 0.56, kw - 0.28, 0.32, l, size=9.5, color=GRIS,
                     align=PP_ALIGN.CENTER, space_after=0)
            y += 1.12
        # columnas o viñetas
        if sl.get("col_izq"):
            for k, key in enumerate(("col_izq", "col_der")):
                if key not in sl:
                    continue
                tit, items = sl[key]
                x = 0.45 + k * (left_w / 2)
                cw = left_w / 2 - 0.3
                _txt(s, x, y, cw, 0.35, tit, size=14, color=AZUL if k == 0 else ROJO, bold=True,
                     space_after=2)
                _rect(s, x, y + 0.36, 1.1, 0.045, ACERO if k == 0 else ROJO)
                _txt(s, x, y + 0.5, cw, 4.6, ["•  " + it for it in items], size=12.5,
                     color=RGBColor(0x12, 0x21, 0x2B), space_after=7)
        elif sl.get("bullets"):
            _txt(s, 0.45, y, left_w - 0.2, 4.9, ["•  " + b for b in sl["bullets"]], size=13.5,
                 color=RGBColor(0x12, 0x21, 0x2B), space_after=8)
        # tabla
        if sl.get("tabla"):
            rows = sl["tabla"]
            nr, nc = len(rows), len(rows[0])
            th = min(0.34 * nr, 7.5 - y - 1.0)
            gt = s.shapes.add_table(nr, nc, Inches(0.45), Inches(y + (0.15 if sl.get("bullets")
                                                                       else 0)),
                                    Inches(left_w - 0.2), Inches(th)).table
            widths = ([0.30] + [0.70 / (nc - 1)] * (nc - 1)) if nc > 1 else [1.0]
            if nc == 3:
                widths = [0.16, 0.56, 0.28]
            if nc == 4:
                widths = [0.16, 0.24, 0.42, 0.18]
            for c, wd in enumerate(widths):
                gt.columns[c].width = Emu(int(Inches((left_w - 0.2) * wd)))
            for r in range(nr):
                for c in range(nc):
                    cell = gt.cell(r, c)
                    cell.text = rows[r][c]
                    cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
                    cell.margin_left = Pt(4); cell.margin_right = Pt(4)
                    p = cell.text_frame.paragraphs[0]
                    p.font.size = Pt(10.5 if nr <= 8 else 9.5)
                    p.font.name = "Calibri"
                    if r == 0:
                        cell.fill.solid(); cell.fill.fore_color.rgb = AZUL
                        p.font.color.rgb = BLANCO; p.font.bold = True
                    else:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = (BLANCO if r % 2 else CLARO)
                        p.font.color.rgb = RGBColor(0x12, 0x21, 0x2B)
        # imágenes columna derecha
        if imgs:
            box_h = 7.5 - y - 1.15
            slot = box_h / len(imgs)
            for k, im in enumerate(imgs):
                _img(s, im, 0.45 + left_w + 0.15, y + k * slot, 12.45 - left_w, slot - 0.12)
        # nota
        if sl.get("nota"):
            nh = 0.62 + 0.115 * (len(sl["nota"]) // 105)
            _rect(s, 0.45, 7.5 - 0.62 - nh + 0.35, 12.45, nh, FONDO_NOTA)
            _rect(s, 0.45, 7.5 - 0.62 - nh + 0.35, 0.07, nh, ROJO)
            _txt(s, 0.62, 7.5 - 0.55 - nh + 0.4, 12.2, nh, sl["nota"], size=10.5,
                 color=RGBColor(0x7D, 0x1F, 0x24), space_after=0)
        # pie
        _rect(s, 0.45, 7.16, 12.45, 0.012, AZUL)
        _txt(s, 0.45, 7.2, 10.6, 0.25, SELLO, size=8, color=ROJO, space_after=0)
        _txt(s, 11.2, 7.2, 1.7, 0.25, f"{i} / {n}", size=10, color=AZUL, bold=True,
             align=PP_ALIGN.RIGHT, space_after=0)
    prs.save(path)


if __name__ == "__main__":
    os.makedirs("entregables", exist_ok=True)
    S2 = deck()
    print("láminas del mazo:", len(S2))
    PR.S = S2
    base = f"entregables/{COD}_PRESENTACION-CONSEJO-UNIVERSITARIO"
    render_pptx(S2, base + ".pptx")
    PR.render_pdf(base + ".pdf")
    PR.render_html(base + ".html")
    for ext in (".pptx", ".pdf", ".html"):
        print(f"  OK {base}{ext}  {os.path.getsize(base+ext)/1024:.0f} kB")
