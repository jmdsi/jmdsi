#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
md2x · Conversor del subconjunto markdown del expediente UMC P04 a HTML / DOCX / PDF.
Subconjunto soportado:
  # / ## / ###      encabezados
  párrafos          texto con **negritas**
  - viñetas
  | tablas |        primera fila = encabezado
  ![cap](ruta)      imagen centrada
  > nota            recuadro de advertencia
  ---               separador (se omite en salida)
PDF: portada + encabezado/pie con sello obligatorio + marcadores de navegación.
"""
import base64
import io
import os
import re

# ------------------------------------------------------------------ parsing
BOLD = re.compile(r"\*\*(.+?)\*\*")


def _inline(s):
    out, pos = [], 0
    for m in BOLD.finditer(s):
        out.append(("t", s[pos:m.start()])); out.append(("b", m.group(1))); pos = m.end()
    out.append(("t", s[pos:]))
    return out


def parse(text):
    blocks, buf, mode = [], [], None
    def flush():
        nonlocal buf, mode
        if not buf:
            mode = None; return
        if mode == "para":
            blocks.append(("p", " ".join(buf)))
        elif mode == "bul":
            blocks.append(("ul", list(buf)))
        elif mode == "note":
            blocks.append(("note", list(buf)))
        buf, mode = [], None
    for line in text.splitlines():
        s = line.rstrip()
        if not s.strip():
            flush(); continue
        if s.startswith("|"):
            flush()
            cells = [c.strip() for c in s.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            if blocks and blocks[-1][0] == "table":
                blocks[-1][1].append(cells)
            else:
                blocks.append(("table", [cells]))
            continue
        if s.startswith("### "):
            flush(); blocks.append(("h3", s[4:])); continue
        if s.startswith("## "):
            flush(); blocks.append(("h2", s[3:])); continue
        if s.startswith("# "):
            flush(); blocks.append(("h1", s[2:])); continue
        if s.startswith("---"):
            flush(); continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s.strip())
        if m:
            flush(); blocks.append(("img", m.group(2), m.group(1))); continue
        if s.startswith("> "):
            flush() if mode not in (None, "note") else None
            mode = "note"; buf.append(s[2:]); continue
        if s.startswith("- "):
            if mode != "bul":
                flush(); mode = "bul"
            buf.append(s[2:]); continue
        if mode != "para":
            flush(); mode = "para"
        buf.append(s.strip())
    flush()
    return blocks


# ------------------------------------------------------------------ HTML
CSS = """
:root{--azul:#0b3c5d;--acero:#1d5b7a;--rojo:#a4262c;--gris:#f4f6f8;--linea:#c9d4dc}
*{box-sizing:border-box}body{font-family:"Segoe UI",Helvetica,Arial,sans-serif;color:#12212b;
 line-height:1.55;margin:0;background:#e8ecef}
nav#toc{position:fixed;left:0;top:0;width:250px;height:100vh;overflow:auto;background:var(--azul);
 color:#dce6ee;padding:18px 14px;font-size:12px}
nav#toc a{color:#cfe0ea;text-decoration:none;display:block;padding:2px 0}
nav#toc a.h2{padding-left:12px;font-size:11px;color:#9db8c9}
main{margin-left:250px;max-width:900px;background:#fff;padding:40px 48px 80px;
 box-shadow:0 2px 16px rgba(0,0,0,.15)}
h1{color:var(--azul);font-size:1.45rem;border-bottom:3px solid var(--azul);padding-bottom:.3em;
 margin-top:2.2em;page-break-after:avoid}
h2{color:var(--azul);font-size:1.12rem;margin-top:1.7em;page-break-after:avoid}
h3{color:var(--acero);font-size:1rem;margin-top:1.4em}
p,li{font-size:.9rem}
table{border-collapse:collapse;width:100%;margin:1em 0;font-size:.8rem;page-break-inside:auto}
th{background:var(--azul);color:#fff;text-align:left;padding:6px 8px;border:1px solid var(--azul)}
td{border:1px solid var(--linea);padding:5px 8px;vertical-align:top}
tbody tr:nth-child(even){background:var(--gris)}
tr{page-break-inside:avoid}
figure{margin:1.2em 0;text-align:center;page-break-inside:avoid}
figure img{max-width:100%;border:1px solid var(--linea)}
figcaption{font-size:.76rem;color:#33475a;margin-top:.4em}
.note{border-left:5px solid var(--rojo);background:#fdf3f3;padding:.7em 1em;margin:1.1em 0;
 font-size:.85rem;white-space:pre-line}
.cover{background:var(--azul);color:#fff;padding:44px 40px;margin:-40px -48px 30px}
.cover h1{color:#fff;border:0;margin:0 0 .3em;font-size:1.9rem}
.cover table{color:#12212b}
.sello{color:#ffd9d6;font-weight:700;font-size:.8rem;margin-top:1.4em}
@media print{@page{size:A4;margin:15mm 13mm}nav#toc{display:none}main{margin:0;box-shadow:none;
 padding:0;max-width:none}.cover{margin:0 0 20mm}}
"""


def _b64(path):
    ext = os.path.splitext(path)[1].lower()
    mime = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg"}.get(ext, "image/png")
    if not os.path.exists(path):
        return None, None
    data = open(path, "rb").read()
    return mime, base64.b64encode(data).decode()


def render_html(blocks, cover_rows, title, embed=True):
    toc, body = [], []
    for b in blocks:
        if b[0] == "h1":
            toc.append(("h1", b[1])); body.append(f"<h1>{b[1]}</h1>")
        elif b[0] == "h2":
            toc.append(("h2", b[1])); body.append(f"<h2>{b[1]}</h2>")
        elif b[0] == "h3":
            body.append(f"<h3>{b[1]}</h3>")
        elif b[0] == "p":
            html = "".join(f"<b>{t}</b>" if k == "b" else t for k, t in _inline(b[1]))
            body.append(f"<p>{html}</p>")
        elif b[0] == "ul":
            items = "".join("<li>" + "".join(f"<b>{t}</b>" if k == "b" else t
                                             for k, t in _inline(x)) + "</li>" for x in b[1])
            body.append(f"<ul>{items}</ul>")
        elif b[0] == "note":
            body.append("<div class='note'>" + "<br>".join(b[1]) + "</div>")
        elif b[0] == "table":
            rows = b[1]
            head = "".join(f"<th>{c}</th>" for c in rows[0])
            rest = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows[1:])
            body.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rest}</tbody></table>")
        elif b[0] == "img":
            path, cap = b[1], b[2]
            if embed:
                mime, b64 = _b64(path)
                src = f"data:{mime};base64,{b64}" if b64 else ""
            else:
                src = path
            body.append(f"<figure><img src='{src}' alt='{cap}'/><figcaption>{cap}</figcaption></figure>")
    nav = "".join(f"<a class='{k}' href='#{i}'>{t}</a>" for i, (k, t) in enumerate(toc))
    # anclas
    idx = 0
    out = []
    for line in body:
        if line.startswith("<h1>"):
            line = line.replace("<h1>", f"<h1 id='{idx}'>", 1); idx += 1
        elif line.startswith("<h2>"):
            line = line.replace("<h2>", f"<h2 id='{idx}'>", 1); idx += 1
        out.append(line)
    cover = "<table>" + "".join(
        f"<tr><th style='width:34%'>{k}</th><td>{v}</td></tr>" for k, v in cover_rows) + "</table>"
    return f"""<!DOCTYPE html><html lang='es'><head><meta charset='utf-8'>
<meta name='viewport' content='width=device-width,initial-scale=1'><title>{title}</title>
<style>{CSS}</style></head><body><nav id='toc'><b>{title}</b><hr style='border-color:#1d5b7a'>{nav}</nav>
<main><div class='cover'><h1>{title}</h1>{cover}
<div class='sello'>INGENIERÍA BÁSICA / ANTEPROYECTO — EMITIDO PARA COORDINACIÓN Y PRESUPUESTO
PRELIMINAR — NO APTO PARA CONSTRUCCIÓN</div></div>
{''.join(out)}</main></body></html>"""


# ------------------------------------------------------------------ DOCX
def render_docx(blocks, cover_rows, title, path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    for nm, sz, col in [("Heading 1", 15, "0B3C5D"), ("Heading 2", 12, "0B3C5D"),
                        ("Heading 3", 11, "1D5B7A")]:
        s = doc.styles[nm]; s.font.size = Pt(sz); s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(col)
    p = doc.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string("0B3C5D")
    for k, v in cover_rows:
        pp = doc.add_paragraph(); rr = pp.add_run(f"{k}: "); rr.bold = True
        pp.add_run(v)
    sp = doc.add_paragraph(); sr = sp.add_run("INGENIERÍA BÁSICA / ANTEPROYECTO — EMITIDO PARA "
                                              "COORDINACIÓN Y PRESUPUESTO PRELIMINAR — NO APTO "
                                              "PARA CONSTRUCCIÓN")
    sr.bold = True; sr.font.color.rgb = RGBColor.from_string("A4262C"); sr.font.size = Pt(9)
    doc.add_page_break()
    for b in blocks:
        if b[0] == "h1":
            doc.add_heading(b[1], 1)
        elif b[0] == "h2":
            doc.add_heading(b[1], 2)
        elif b[0] == "h3":
            doc.add_heading(b[1], 3)
        elif b[0] == "p":
            _docx_para(doc, b[1])
        elif b[0] == "ul":
            for x in b[1]:
                pp = doc.add_paragraph(style="List Bullet"); _docx_runs(pp, x)
        elif b[0] == "note":
            t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
            c = t.cell(0, 0); c.text = ""
            rr = c.paragraphs[0].add_run("\n".join(b[1]))
            rr.font.size = Pt(9); rr.bold = True
            rr.font.color.rgb = RGBColor.from_string("A4262C")
        elif b[0] == "table":
            rows = b[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    cell = t.cell(i, j); cell.text = ""
                    rr = cell.paragraphs[0].add_run(c); rr.font.size = Pt(8)
                    if i == 0:
                        rr.bold = True
        elif b[0] == "img":
            if os.path.exists(b[1]):
                doc.add_picture(b[1], width=Inches(6.3))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(b[2]); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].font.size = Pt(8)
    doc.save(path)


def _docx_runs(pp, s):
    for k, t in _inline(s):
        r = pp.add_run(t); r.bold = (k == "b")


def _docx_para(doc, s):
    pp = doc.add_paragraph(); _docx_runs(pp, s)


# ------------------------------------------------------------------ PDF
def _render_pdf_pass(blocks, cover_rows, title, path, autor, TOTAL):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
                                    Table, TableStyle, Image, PageBreak, Flowable)
    from reportlab.pdfbase import pdfmetrics

    SELLO = ("INGENIERÍA BÁSICA / ANTEPROYECTO — EMITIDO PARA COORDINACIÓN Y PRESUPUESTO "
             "PRELIMINAR — NO APTO PARA CONSTRUCCIÓN")
    HEAD = ("UNIVERSIDAD MARÍTIMA DEL CARIBE – PLAN MAESTRO UMC 2025–2030  |  TORRE DE "
            "OBSERVACIÓN / PUENTE DE MANDO Y SALA TÉCNICA DE INVESTIGACIÓN HIDROMETEOROLÓGICA")
    W, H = A4
    ss = getSampleStyleSheet()
    s_h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=13.5, leading=17,
                          textColor=colors.HexColor("#0b3c5d"), spaceBefore=14, spaceAfter=7,
                          borderWidth=0)
    s_h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=11, leading=14,
                          textColor=colors.HexColor("#0b3c5d"), spaceBefore=10, spaceAfter=5)
    s_h3 = ParagraphStyle("h3", parent=ss["Heading3"], fontSize=9.8, leading=12.5,
                          textColor=colors.HexColor("#1d5b7a"), spaceBefore=8, spaceAfter=4)
    s_p = ParagraphStyle("p", parent=ss["BodyText"], fontSize=8.7, leading=12.2,
                         alignment=TA_JUSTIFY, spaceAfter=5)
    s_li = ParagraphStyle("li", parent=s_p, leftIndent=10, bulletIndent=2, spaceAfter=2.5)
    s_note = ParagraphStyle("note", parent=s_p, fontSize=8.2, leading=11.4,
                            textColor=colors.HexColor("#7d1f24"), backColor=colors.HexColor("#fdf3f3"),
                            borderColor=colors.HexColor("#a4262c"), borderWidth=1.4,
                            borderPadding=6, spaceBefore=6, spaceAfter=8)
    s_cap = ParagraphStyle("cap", parent=s_p, fontSize=7.4, leading=9.6,
                           textColor=colors.HexColor("#33475a"), alignment=1, spaceBefore=2)
    s_tb = ParagraphStyle("tb", parent=s_p, fontSize=7.4, leading=9.4, spaceAfter=0,
                          alignment=0)
    s_tbh = ParagraphStyle("tbh", parent=s_tb, textColor=colors.white, fontSize=7.4)

    class BM(Flowable):
        def __init__(self, key, txt, lvl):
            Flowable.__init__(self); self.key, self.txt, self.lvl = key, txt, lvl
            self.width = self.height = 0
        def draw(self):
            pass  # el marcador se registra en Doc.afterFlowable (página correcta)

    from reportlab.lib.utils import ImageReader

    def _imgsize(pth):
        return ImageReader(pth).getSize()

    story, bm = [], 0
    for b in blocks:
        if b[0] == "h1":
            bm += 1; story.append(BM(f"bm{bm}", b[1], 0)); story.append(Paragraph(b[1], s_h1))
        elif b[0] == "h2":
            bm += 1; story.append(BM(f"bm{bm}", b[1], 1)); story.append(Paragraph(b[1], s_h2))
        elif b[0] == "h3":
            story.append(Paragraph(b[1], s_h3))
        elif b[0] == "p":
            story.append(Paragraph(_rl(b[1]), s_p))
        elif b[0] == "ul":
            for x in b[1]:
                story.append(Paragraph(_rl(x), s_li, bulletText="•"))
        elif b[0] == "note":
            story.append(Paragraph("<br/>".join(_rl(x) for x in b[1]), s_note))
        elif b[0] == "table":
            rows = b[1]
            data = [[Paragraph(_rl(c), s_tbh if i == 0 else s_tb) for c in r]
                    for i, r in enumerate(rows)]
            t = Table(data, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b3c5d")),
                ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#c9d4dc")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f4f6")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(Spacer(1, 4)); story.append(t); story.append(Spacer(1, 6))
        elif b[0] == "img":
            pth = b[1]
            if os.path.exists(pth):
                try:
                    iw, ih = _imgsize(pth)
                    maxw = W - 40 * mm
                    maxh = H - 70 * mm
                    sc = min(maxw / iw, maxh / ih)
                    story.append(Image(pth, width=iw * sc, height=ih * sc))
                    story.append(Paragraph(b[2], s_cap))
                except Exception:
                    pass
    # portada
    cover = [Paragraph("<b>REPÚBLICA BOLIVARIANA DE VENEZUELA</b><br/>"
                       "MINISTERIO DEL PODER POPULAR PARA EL TRANSPORTE ACUÁTICO Y AÉREO<br/>"
                       "MARINA MERCANTE DE VENEZUELA<br/>"
                       "<b>UNIVERSIDAD MARÍTIMA DEL CARIBE – UMC</b>", ss["BodyText"]),
             Spacer(1, 10),
             Paragraph(f"<b><font size=17 color='#0b3c5d'>{title}</font></b>",
                       ParagraphStyle("port", parent=ss["BodyText"], fontSize=17, leading=22)),
             Spacer(1, 12)]
    data = [[Paragraph(f"<b>{k}</b>", s_tb), Paragraph(v, s_tb)] for k, v in cover_rows]
    tc = Table(data, colWidths=[52 * mm, 118 * mm], hAlign="LEFT")
    tc.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#0b3c5d")),
                            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef2f5")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 3),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    cover += [tc, Spacer(1, 14),
              Paragraph(f"<b><font color='#a4262c'>{SELLO}</font></b>", ss["BodyText"]),
              Spacer(1, 6),
              Paragraph("Este documento constituye INGENIERÍA BÁSICA / ANTEPROYECTO. No sustituye "
                        "el levantamiento topográfico, el estudio geotécnico, la inspección "
                        "estructural post-sísmica, el cálculo estructural, las memorias de "
                        "instalaciones, los permisos ni los planos firmados por los profesionales "
                        "responsables. NO DEBERÁ UTILIZARSE PARA REPLANTEO O CONSTRUCCIÓN.",
                        ParagraphStyle("adv", parent=s_p, fontSize=8, textColor=colors.HexColor("#33475a"))),
              PageBreak()]

    class Doc(BaseDocTemplate):
        def __init__(self, fn, **kw):
            BaseDocTemplate.__init__(self, fn, **kw)
            fr = Frame(18 * mm, 16 * mm, W - 36 * mm, H - 40 * mm, id="n")
            self.addPageTemplates([PageTemplate(id="all", frames=[fr], onPage=self._deco)])

        def afterFlowable(self, flowable):
            if isinstance(flowable, BM):
                self.canv.bookmarkPage(flowable.key)
                self.canv.addOutlineEntry(flowable.txt[:120], flowable.key,
                                          level=flowable.lvl, closed=0)
        def _deco(self, c, d):
            c.saveState()
            c.setFont("Helvetica-Bold", 6.6); c.setFillColor(colors.HexColor("#0b3c5d"))
            c.drawString(18 * mm, H - 12 * mm, HEAD)
            c.setFont("Helvetica-Bold", 6.2); c.setFillColor(colors.HexColor("#a4262c"))
            c.drawString(18 * mm, H - 15.4 * mm, SELLO)
            c.setStrokeColor(colors.HexColor("#0b3c5d")); c.setLineWidth(.7)
            c.line(18 * mm, H - 17 * mm, W - 18 * mm, H - 17 * mm)
            c.line(18 * mm, 13 * mm, W - 18 * mm, 13 * mm)
            c.setFont("Helvetica", 6.6); c.setFillColor(colors.HexColor("#33475a"))
            c.drawString(18 * mm, 9.6 * mm, "UMC2030-P04-TOBSHM-R02 · 17 de octubre de 2026")
            pg = c.getPageNumber()
            c.drawRightString(W - 18 * mm, 9.6 * mm,
                              f"Pág. {pg} de {TOTAL}" if TOTAL else f"Pág. {pg}")
            c.restoreState()

    from reportlab.pdfgen import canvas as _canvas

    class NumCanvas(_canvas.Canvas):
        def __init__(self, *a, **kw):
            _canvas.Canvas.__init__(self, *a, **kw); self._saved = []
        def showPage(self):
            self._saved.append(dict(self.__dict__)); self._startPage()
        def save(self):
            total = len(self._saved)
            for st in self._saved:
                self.__dict__.update(st)
                self.saveState()
                self.setFont("Helvetica", 6.6)
                self.setFillColor(colors.HexColor("#33475a"))
                self.drawRightString(W - 18 * mm, 9.6 * mm,
                                     f"Pág. {self._pageNumber} de {total}")
                self.restoreState()
                _canvas.Canvas.showPage(self)
            _canvas.Canvas.save(self)

    doc = Doc(path, pagesize=A4, title=title, author=autor,
              subject="Plan Maestro UMC 2025-2030 · Proyecto de Planta Física e Infraestructura",
              keywords="UMC; Plan Maestro 2030; torre de observación; hidrometeorología; patrimonio")
    doc.multiBuild(cover + story)


def render_pdf(blocks, cover_rows, title, path,
               autor="Universidad Marítima del Caribe – UMC"):
    """Dos pasadas: la primera determina el total de páginas que imprime el pie."""
    tmp = path + ".tmp"
    _render_pdf_pass(blocks, cover_rows, title, tmp, autor, None)
    import pymupdf
    d = pymupdf.open(tmp)
    n = d.page_count
    d.close()
    os.remove(tmp)
    _render_pdf_pass(blocks, cover_rows, title, path, autor, n)


def _rl(s):
    """Escapa HTML y convierte **negrita** a <b>."""
    return _rl2(s)


def _rl2(s):
    parts = s.split("**")
    out = []
    for i, p in enumerate(parts):
        p = p.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        out.append(f"<b>{p}</b>" if i % 2 else p)
    return "".join(out)
