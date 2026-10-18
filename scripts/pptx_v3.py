#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""E6 v3 · Parte PPTX: toma el mazo editado por el Rectorado en Drive
(«AGREGAR NOTAS EXPLICATIVAS Y ABREVIATURAS…PPTX», 31 láminas con sus aportes) y le aplica,
SIN reconstruirlo, los tres pedidos:
  1. Notas explicativas del expositor en las 32 láminas (notes pane).
  2. Lámina nueva «Siglas y abreviaturas de la exposición» en posición 3.
  3. Lámina ROM re-hecha: G-13 etiquetada por capítulo + cuadro «Cómo leer la figura» +
     KPI reubicados + nota de lectura; agenda re-numerada y pies «n / 32» corregidos."""
import os
import re

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

import notas_v3 as NV

SRC = ("google_drive/AGREGAR NOTAS EXPLICATIVAS Y ABREVIATURAS DE LA EXPOSICION "
       "UMC2030-P04-TOBSHM-R02_PRESENTACION-CONSEJO-UNIVERSITARIO.pptx")
DST = "entregables/UMC2030-P04-TOBSHM-R02_PRESENTACION-CONSEJO-UNIVERSITARIO.pptx"
G13 = "build/fig/G-13_rom.png"
LOGO = "src/assets/user/UMC-EMBLEMA-172.png"

AZUL = RGBColor(0x0B, 0x3C, 0x5D)
ACERO = RGBColor(0x1D, 0x5B, 0x7A)
ROJO = RGBColor(0xA4, 0x26, 0x2C)
GRIS = RGBColor(0x33, 0x47, 0x5A)
CLARO = RGBColor(0xEE, 0xF2, 0xF5)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONDO_NOTA = RGBColor(0xFD, 0xF3, 0xF3)
TINTA = RGBColor(0x12, 0x21, 0x2B)

SELLO = ("INGENIERÍA BÁSICA / ANTEPROYECTO — EMITIDO PARA COORDINACIÓN Y PRESUPUESTO "
         "PRELIMINAR — NO APTO PARA CONSTRUCCIÓN")
IN = Inches


def _txt(slide, l, t, w, h, runs, size=14, color=GRIS, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=6, line=1.12):
    box = slide.shapes.add_textbox(IN(l), IN(t), IN(w), IN(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [runs]
    for i, r in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line
        run = p.add_run()
        run.text = r
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = "Calibri"
    return box


def _rect(slide, l, t, w, h, fill):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, IN(l), IN(t), IN(w), IN(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _drop(shape):
    sp = shape._element
    sp.getparent().remove(sp)


def slide_title(sl):
    best = ""
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip().splitlines()[0]
            if len(t) > len(best) and sh.top is not None and IN(0.10) <= sh.top <= IN(0.80):
                best = t
    return best


def add_glossary(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, 13.333, 1.32, AZUL)
    _rect(s, 0, 1.32, 13.333, 0.075, ROJO)
    _txt(s, 0.45, 0.22, 10.4, 0.75, "Siglas y abreviaturas de la exposición", size=25,
         color=BLANCO, bold=True, space_after=2)
    _txt(s, 0.47, 0.86, 10.4, 0.4,
         "Guía de lectura rápida: todos los términos que aparecen en esta exposición",
         size=12, color=RGBColor(0xCF, 0xE0, 0xEA), space_after=0)
    if os.path.exists(LOGO):
        s.shapes.add_picture(LOGO, IN(10.97), IN(0.03), IN(1.28), IN(1.28))
    _txt(s, 0.45, 1.60, 6.05, 4.6, ["•  " + b for b in NV.GLOSARIO_IZQ], size=10.2,
         color=TINTA, space_after=5)
    _txt(s, 6.85, 1.60, 6.05, 4.6, ["•  " + b for b in NV.GLOSARIO_DER], size=10.2,
         color=TINTA, space_after=5)
    _rect(s, 0.45, 6.44, 12.45, 0.62, FONDO_NOTA)
    _rect(s, 0.45, 6.44, 0.07, 0.62, ROJO)
    _txt(s, 0.62, 6.55, 12.2, 0.5,
         "Fuente: anexo C del expediente R02 (§27.3), ampliado con los términos propios de "
         "esta exposición. Definiciones completas en el glosario del documento.",
         size=10.5, color=RGBColor(0x7D, 0x1F, 0x24), space_after=0)
    _rect(s, 0.45, 7.16, 12.45, 0.012, AZUL)
    _txt(s, 0.45, 7.20, 10.6, 0.25, SELLO, size=8, color=ROJO, space_after=0)
    _txt(s, 11.2, 7.20, 1.7, 0.25, "3 / 32", size=10, color=AZUL, bold=True,
         align=PP_ALIGN.RIGHT, space_after=0)
    # mover a la posición 3
    lst = prs.slides._sldIdLst
    el = list(lst)[-1]
    lst.remove(el)
    lst.insert(2, el)
    return s


def fix_rom(prs):
    sl = None
    for s in prs.slides:
        if slide_title(s).startswith("Inversión estimada"):
            sl = s
            break
    assert sl is not None, "no se encontró la lámina ROM"
    # 1) fuera la gráfica vieja (imagen grande desplazada por el usuario)
    for sh in list(sl.shapes):
        if sh.shape_type == 13 and sh.width > IN(7.0) and sh.top > IN(1.4) and sh.left < IN(1):
            _drop(sh)
    # 2) KPI: reubicar en rejilla 2×2 a la derecha
    rects = [sh for sh in sl.shapes if sh.shape_type == 1
             and sh.top is not None and IN(1.50) <= sh.top <= IN(1.72)
             and IN(0.80) <= sh.height <= IN(1.00)]
    order = sorted(range(len(rects)), key=lambda k: rects[k].left)
    grid = [(8.05, 1.58), (10.55, 1.58), (8.05, 2.66), (10.55, 2.66)]
    for rank, idx in enumerate(order):
        r = rects[idx]
        ox = r.left
        x, y = grid[rank]
        r.left, r.top, r.width, r.height = IN(x), IN(y), IN(2.35), IN(0.95)
        for sh in sl.shapes:
            if sh is r or sh.left is None:
                continue
            d = sh.left - ox
            if IN(-0.02) <= d <= IN(0.12):
                if sh.shape_type == 1 and sh.height < IN(0.12):          # filete de color
                    sh.left, sh.top, sh.width, sh.height = IN(x), IN(y), IN(2.35), IN(0.06)
                elif sh.has_text_frame and IN(1.65) <= sh.top <= IN(1.85):  # valor
                    sh.left, sh.top, sh.width, sh.height = IN(x + 0.06), IN(y + 0.10), IN(2.23), IN(0.45)
                elif sh.has_text_frame and IN(2.05) <= sh.top <= IN(2.35):  # etiqueta
                    sh.left, sh.top, sh.width, sh.height = IN(x + 0.06), IN(y + 0.55), IN(2.23), IN(0.38)
    # 3) nueva G-13 etiquetada
    sl.shapes.add_picture(G13, IN(0.30), IN(1.56), width=IN(7.35))
    # 4) cuadro «Cómo leer la figura»
    _rect(sl, 7.98, 3.86, 4.92, 2.34, BLANCO)
    _rect(sl, 7.98, 3.86, 0.07, 2.34, ACERO)
    box = sl.shapes.add_textbox(IN(8.16), IN(3.94), IN(4.62), IN(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.space_after = Pt(4)
    r = p.add_run()
    r.text = "Cómo leer la figura G-13"
    r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = AZUL; r.font.name = "Calibri"
    for b in NV.ROM_GUIA:
        p = tf.add_paragraph()
        p.space_after = Pt(3)
        p.line_spacing = 1.05
        r = p.add_run()
        r.text = "•  " + b
        r.font.size = Pt(8.8); r.font.color.rgb = TINTA; r.font.name = "Calibri"
    # 5) nota inferior: geometría + texto de lectura
    for sh in sl.shapes:
        if sh.shape_type == 1 and sh.top is not None and IN(6.30) <= sh.top <= IN(6.60) \
                and sh.width > IN(10):
            sh.top, sh.height = IN(6.28), IN(0.94)
        if sh.shape_type == 1 and sh.top is not None and IN(6.30) <= sh.top <= IN(6.60) \
                and sh.width < IN(0.2):
            sh.top, sh.height = IN(6.28), IN(0.94)
        if sh.has_text_frame and sh.text_frame.text.strip().startswith(
                "No es presupuesto de licitación"):
            sh.top, sh.height = IN(6.34), IN(0.86)
            tf = sh.text_frame
            tf.clear()
            r = tf.paragraphs[0].add_run()
            r.text = NV.ROM_NOTA
            r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x7D, 0x1F, 0x24)
            r.font.name = "Calibri"
    return sl


def fix_agenda(prs):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_table and sh.table.cell(0, 1).text.strip().lower().startswith("bloque"):
                tb = sh.table
                for i, row in enumerate(NV.AGENDA, start=1):
                    for c, val in enumerate(row):
                        tb.cell(i, c).text = val
                return True
    return False


def main():
    prs = Presentation(SRC)
    add_glossary(prs)
    fix_rom(prs)
    ok_ag = fix_agenda(prs)
    n = len(prs.slides.__iter__.__self__._sldIdLst)
    assert n == 32, f"esperaba 32 láminas, hay {n}"
    assert len(NV.NOTES) == n
    # notas del expositor
    for i, s in enumerate(prs.slides, 1):
        s.notes_slide.notes_text_frame.text = (
            f"[{i}/{n}] " + NV.NOTES[i - 1])
    # renumeración de pies
    pat = re.compile(r"^\s*\d+\s*/\s*\d+\s*$")
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and pat.match(sh.text_frame.text or ""):
                tf = sh.text_frame
                new = f"{i} / {n}"
                if tf.paragraphs[0].runs:
                    tf.paragraphs[0].runs[0].text = new
                    for r in tf.paragraphs[0].runs[1:]:
                        r.text = ""
                    for p in tf.paragraphs[1:]:
                        for r in p.runs:
                            r.text = ""
                else:
                    r = tf.paragraphs[0].add_run()
                    r.text = new
                    r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = AZUL
                    r.font.name = "Calibri"
    os.makedirs("entregables", exist_ok=True)
    prs.save(DST)
    print(f"OK {DST} · {n} láminas · agenda={ok_ag} · "
          f"{os.path.getsize(DST)/1e6:.2f} MB")


if __name__ == "__main__":
    main()
